#!/usr/bin/env python3
"""Build Word (.docx) files for the book from Markdown sources, using python-docx.

Usage
-----
    python build/build_docx.py book/ch03.md -o build/output/ch03.docx
    python build/build_docx.py book/00_preface.md book/00_outline.md -o build/output/00_front.docx
    python build/build_docx.py ... --pdf          # also export a PDF with LibreOffice

Supported Markdown (see STYLE_GUIDE.md, section 8 for the full list)
--------------------------------------------------------------------
  front matter       ---\ntitle: ...\nsubtitle: ...\nlabel: ...\n---   (title page)
  headings           #, ##, ###, ####   (# starts a new page)
  paragraphs         **bold**, *italic*, `code`, [link](https://...)
  lists              - bullet, 1. numbered, - [ ] checklist item (2-space indent = level 2)
  tables             GitHub pipe tables
  code               ```python title=code/ch03/size.py   (also pine, sql, bash, json, text)
  include            @include(code/ch03/size.py, python)  -> code box filled from the tested file
  prompt             ```prompt Binance MCP                 -> blue prompt box with a label
  callouts           > [!RISK] / [!NOTE] / [!TIP] / [!UNVERIFIED] / [!CHECK]  optional title
  figures            ![Figure 3.2. Caption.](figures/ch03/fig03-02-name.png){width=14}
  page break         \\pagebreak
  comments           <!-- ignored -->
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

# ---------------------------------------------------------------------------
# Design tokens (keep in sync with STYLE_GUIDE.md)
# ---------------------------------------------------------------------------
BODY_FONT = "Cambria"
HEAD_FONT = "Calibri"
CODE_FONT = "Consolas"
INK = "0B0B0B"
INK2 = "52514E"
MUTED = "898781"
H1_COLOR = "14213D"
H2_COLOR = "1C5CAB"
HAIRLINE = "C3C2B7"
CODE_FILL = "F4F4F2"
PROMPT_FILL = "EEF4FC"
PROMPT_EDGE = "2A78D6"
TABLE_HEAD_FILL = "DCE8F7"
CALLOUTS = {  # type: (label, edge colour, fill colour)
    "RISK": ("Risk", "D03B3B", "FBECEC"),
    "NOTE": ("Note", "898781", "F4F4F2"),
    "TIP": ("Tip", "0CA30C", "EAF6EA"),
    "UNVERIFIED": ("Not verified", "FAB219", "FFF6E0"),
    "CHECK": ("Checklist", "2A78D6", "EEF4FC"),
}
TEXT_WIDTH_CM = 16.0
BOOK_SHORT_TITLE = "Trading Crypto Futures with AI Agents"


# ---------------------------------------------------------------------------
# Low-level XML helpers
# ---------------------------------------------------------------------------
# Schema order matters: Word and LibreOffice reject children in the wrong order.
PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                  "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
                  "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
                  "w:rPr", "w:sectPr", "w:pPrChange")
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                       "w:tblDescription", "w:tblPrChange")
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                  "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.insert_element_before(shd, *TCPR_AFTER_SHD)


def _ppr_border(ppr, sides: dict[str, tuple[int, str, int]]) -> None:
    """Add paragraph borders. sides: {'left': (size_eighths_pt, colour, space_pt)}."""
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(pbdr, *PPR_AFTER_PBDR)
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            size, color, space = sides[side]
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), str(space))
            el.set(qn("w:color"), color)
            pbdr.append(el)


def _ppr_shading(ppr, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.insert_element_before(shd, *PPR_AFTER_PBDR[1:])


def _add_field(run, instr: str) -> None:
    """Insert a Word field (e.g. PAGE, TOC) into a run."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, text, sep, placeholder, end):
        run._r.append(el)


def _east_asia(style_or_run_element, name: str) -> None:
    """Set one font for every script slot (ascii, hAnsi, cs, eastAsia) on a style or run."""
    rpr = style_or_run_element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


# ---------------------------------------------------------------------------
# Document setup: page, styles, header/footer
# ---------------------------------------------------------------------------
def _para_style(doc, name, base="Normal", font=BODY_FONT, size=11, color=INK, bold=False,
                italic=False, before=0, after=6, line=1.15, keep_next=False):
    styles = doc.styles
    st = styles[name] if name in [s.name for s in styles] else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if st.type != WD_STYLE_TYPE.PARAGRAPH:
        raise ValueError(f"Style name clash: '{name}' already exists as a non-paragraph style")
    if name != "Normal":
        st.base_style = styles[base]
    st.font.name = font
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.italic = italic
    st.font.color.rgb = RGBColor.from_string(color)
    _east_asia(st.element, font)
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    return st


def setup_document(doc) -> None:
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # the default template omits a required attribute
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.0)

    _para_style(doc, "Normal", font=BODY_FONT, size=11, after=6, line=1.2)
    h1 = _para_style(doc, "Heading 1", font=HEAD_FONT, size=24, color=H1_COLOR, bold=True,
                     before=0, after=14, line=1.05, keep_next=True)
    h1.paragraph_format.page_break_before = True
    _para_style(doc, "Heading 2", font=HEAD_FONT, size=15, color=H2_COLOR, bold=True,
                before=16, after=6, line=1.1, keep_next=True)
    _para_style(doc, "Heading 3", font=HEAD_FONT, size=12.5, color=INK, bold=True,
                before=12, after=4, line=1.1, keep_next=True)
    _para_style(doc, "Heading 4", font=HEAD_FONT, size=11, color=INK2, bold=True,
                before=8, after=3, line=1.1, keep_next=True)
    for lvl in (1, 2, 3, 4):  # remove theme italics/colours some templates add
        doc.styles[f"Heading {lvl}"].font.italic = False

    code = _para_style(doc, "Code Block", font=CODE_FONT, size=8.5, color=INK, after=0, line=1.0)
    ppr = code.element.get_or_add_pPr()
    _ppr_border(ppr, {s: (4, HAIRLINE, 4) for s in ("top", "left", "bottom", "right")})
    _ppr_shading(ppr, CODE_FILL)
    code.paragraph_format.left_indent = Cm(0.15)
    code.paragraph_format.right_indent = Cm(0.15)

    _para_style(doc, "Code Label", font=HEAD_FONT, size=8, color=INK2, bold=True, before=8,
                after=2, keep_next=True)

    prompt = _para_style(doc, "Prompt Text", font=HEAD_FONT, size=10.5, color=INK, after=0, line=1.15)
    ppr = prompt.element.get_or_add_pPr()
    _ppr_border(ppr, {"left": (36, PROMPT_EDGE, 8), "top": (4, PROMPT_FILL, 4),
                      "bottom": (4, PROMPT_FILL, 4), "right": (4, PROMPT_FILL, 4)})
    _ppr_shading(ppr, PROMPT_FILL)
    prompt.paragraph_format.left_indent = Cm(0.35)
    prompt.paragraph_format.right_indent = Cm(0.15)
    _para_style(doc, "Prompt Label", base="Prompt Text", font=HEAD_FONT, size=8.5,
                color=H2_COLOR, bold=True, before=10, after=0, keep_next=True)

    for key, (label, edge, fill) in CALLOUTS.items():
        st = _para_style(doc, f"Callout {key}", font=BODY_FONT, size=10, color=INK, after=0, line=1.15)
        ppr = st.element.get_or_add_pPr()
        _ppr_border(ppr, {"left": (24, edge, 8), "top": (4, fill, 4), "bottom": (4, fill, 4),
                          "right": (4, fill, 4)})
        _ppr_shading(ppr, fill)
        st.paragraph_format.left_indent = Cm(0.35)
        st.paragraph_format.right_indent = Cm(0.15)

    _para_style(doc, "Figure", font=BODY_FONT, size=10, before=10, after=2, keep_next=True)
    _para_style(doc, "Figure Caption", font=HEAD_FONT, size=9, color=INK2, after=12, line=1.1)
    _para_style(doc, "Table Caption", font=HEAD_FONT, size=9, color=INK2, before=8, after=3,
                keep_next=True)
    _para_style(doc, "Table Text", font=HEAD_FONT, size=9.5, after=0, line=1.05)
    _para_style(doc, "Quote Text", font=BODY_FONT, size=10.5, color=INK2, italic=True, after=6)
    _para_style(doc, "Checklist", font=BODY_FONT, size=11, after=3)
    doc.styles["Checklist"].paragraph_format.left_indent = Cm(0.9)
    doc.styles["Checklist"].paragraph_format.first_line_indent = Cm(-0.6)
    _para_style(doc, "Cover Label", font=HEAD_FONT, size=12, color=H2_COLOR, bold=True, after=6)
    _para_style(doc, "Cover Title", font=HEAD_FONT, size=30, color=H1_COLOR, bold=True, after=10,
                line=1.0)
    _para_style(doc, "Cover Subtitle", font=HEAD_FONT, size=14, color=INK2, after=6, line=1.15)
    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.15

    # Hyperlink character style
    if "Hyperlink" not in [s.name for s in doc.styles]:
        hl = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        hl.font.color.rgb = RGBColor.from_string("1C5CAB")
        hl.font.underline = True


def setup_header_footer(section, header_text: str) -> None:
    section.different_first_page_header_footer = True
    hp = section.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run(header_text)
    r.font.name = HEAD_FONT
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    r.font.name = HEAD_FONT
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(INK2)
    _add_field(r, "PAGE")


# ---------------------------------------------------------------------------
# Inline Markdown
# ---------------------------------------------------------------------------
INLINE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<link>\[[^\]]+\]\([^)\s]+\))"
    r"|(?P<bold>\*\*.+?\*\*)"
    r"|(?P<ital>\*[^*\s][^*]*?\*)"
)
ESC_STAR = ""


def _add_hyperlink(paragraph, text: str, url: str, size=None):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), "Hyperlink")
    rpr.append(rstyle)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hl.append(run)
    paragraph._p.append(hl)


def add_inline(paragraph, text: str, bold=False, italic=False, size=None, code_size=None) -> None:
    """Add Markdown-formatted inline text to a paragraph."""
    text = text.replace(r"\*", ESC_STAR)
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _plain(paragraph, text[pos:m.start()], bold, italic, size)
        tok = m.group(0)
        if m.group("code"):
            r = paragraph.add_run(tok[1:-1].replace(ESC_STAR, "*"))
            r.font.name = CODE_FONT
            _east_asia(r._r, CODE_FONT)
            r.font.size = Pt(code_size or ((size or 11) - 1.5))
            r.bold = bold
            r.italic = italic
        elif m.group("link"):
            lm = re.match(r"\[([^\]]+)\]\(([^)\s]+)\)", tok)
            _add_hyperlink(paragraph, lm.group(1).replace(ESC_STAR, "*"), lm.group(2), size)
        elif m.group("bold"):
            add_inline(paragraph, tok[2:-2], True, italic, size, code_size)
        elif m.group("ital"):
            add_inline(paragraph, tok[1:-1], bold, True, size, code_size)
        pos = m.end()
    if pos < len(text):
        _plain(paragraph, text[pos:], bold, italic, size)


def _plain(paragraph, s, bold, italic, size):
    r = paragraph.add_run(s.replace(ESC_STAR, "*"))
    r.bold = bold or None
    r.italic = italic or None
    if size:
        r.font.size = Pt(size)


# ---------------------------------------------------------------------------
# Block builders
# ---------------------------------------------------------------------------
class Builder:
    def __init__(self, doc, src_dir: Path):
        self.doc = doc
        self.src_dir = src_dir
        self.figures: list[str] = []
        self.text_refs: str = ""
        self.warnings: list[str] = []
        self._num_cache: dict[str, int] = {}
        self.first_h1 = True

    # -- numbering restart for ordered lists ---------------------------------
    def _new_list_num(self, style_name: str) -> int:
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        st = self.doc.styles[style_name]
        num_id = int(st.element.pPr.numPr.numId.val)
        abstract = None
        for num in numbering.findall(qn("w:num")):
            if int(num.get(qn("w:numId"))) == num_id:
                abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), abstract)
        num.append(an)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        numbering.append(num)
        return new_id

    def _apply_num(self, paragraph, num_id: int):
        numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()  # inserted in schema order
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = num_id

    # -- blocks ---------------------------------------------------------------
    def title_page(self, meta: dict) -> None:
        d = self.doc
        for _ in range(6):
            d.add_paragraph()
        if meta.get("label"):
            d.add_paragraph(meta["label"], style="Cover Label")
        d.add_paragraph(meta.get("title", ""), style="Cover Title")
        if meta.get("subtitle"):
            d.add_paragraph(meta["subtitle"], style="Cover Subtitle")
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        _ppr_border(p._p.get_or_add_pPr(), {"top": (12, H2_COLOR, 6)})
        if meta.get("author"):
            d.add_paragraph(meta["author"], style="Cover Subtitle")
        if meta.get("note"):
            add_inline(d.add_paragraph(style="Cover Subtitle"), meta["note"], size=10)
        self.first_h1 = False  # title page already started; H1 will break the page

    def heading(self, level: int, text: str) -> None:
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        if level == 1 and self.first_h1:
            p.paragraph_format.page_break_before = False
            self.first_h1 = False
        add_inline(p, text)

    def paragraph(self, text: str) -> None:
        p = self.doc.add_paragraph()
        if re.fullmatch(r"\*\*[^*]+\*\*", text.strip()):
            p.paragraph_format.keep_with_next = True  # a bold lead-in stays with what follows
            p.paragraph_format.space_before = Pt(6)
        add_inline(p, text)
        self.text_refs += " " + text

    def rule(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        _ppr_border(p._p.get_or_add_pPr(), {"bottom": (6, HAIRLINE, 1)})

    def bullet(self, text: str, level: int, ordered: bool, num_id: int | None) -> None:
        if ordered:
            style = "List Number" if level == 0 else "List Number 2"
        else:
            style = "List Bullet" if level == 0 else "List Bullet 2"
        p = self.doc.add_paragraph(style=style)
        if ordered and level == 0 and num_id:
            self._apply_num(p, num_id)
        add_inline(p, text)
        self.text_refs += " " + text

    def checklist(self, text: str, checked: bool) -> None:
        p = self.doc.add_paragraph(style="Checklist")
        r = p.add_run(("☑" if checked else "☐") + "  ")
        r.font.name = "Segoe UI Symbol"
        add_inline(p, text)
        self.text_refs += " " + text

    def code(self, lines: list[str], lang: str, title: str | None) -> None:
        label = (lang or "text").upper()
        if title:
            label += f"  ·  {title}"
        self.doc.add_paragraph(label, style="Code Label")
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph(style="Code Block")
            if i < len(lines) - 1:
                p.paragraph_format.keep_with_next = len(lines) <= 25
            r = p.add_run(line.replace("\t", "    ") if line else " ")
            r.font.name = CODE_FONT
            _east_asia(r._r, CODE_FONT)
        # breathing space after the box
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        sp.paragraph_format.line_spacing = 0.6

    def prompt(self, lines: list[str], label: str | None) -> None:
        head = "PROMPT" + (f"  ·  {label}" if label else "")
        self.doc.add_paragraph(head, style="Prompt Label")
        # merge wrapped lines into paragraphs; blank line = new paragraph
        paras, cur = [], []
        for ln in lines:
            if ln.strip() == "":
                if cur:
                    paras.append(cur)
                    cur = []
                else:
                    paras.append([])
            else:
                cur.append(ln)
        if cur:
            paras.append(cur)
        for k, chunk in enumerate(paras):
            p = self.doc.add_paragraph(style="Prompt Text")
            p.paragraph_format.keep_with_next = k < len(paras) - 1 and len(paras) < 14
            for j, ln in enumerate(chunk):
                if j:
                    p.add_run().add_break(WD_BREAK.LINE)
                self._prompt_line(p, ln)
            if not chunk:
                p.add_run(" ")
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        sp.paragraph_format.line_spacing = 0.6

    @staticmethod
    def _prompt_line(p, ln: str) -> None:
        """Prompt text is plain; {PLACEHOLDERS} are bold so the reader knows to replace them."""
        pos = 0
        for m in re.finditer(r"\{[A-Z0-9_]+\}", ln):
            if m.start() > pos:
                p.add_run(ln[pos:m.start()])
            r = p.add_run(m.group(0))
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(H2_COLOR)
            pos = m.end()
        if pos < len(ln):
            p.add_run(ln[pos:])

    def callout(self, kind: str, title: str | None, lines: list[str]) -> None:
        kind = kind.upper()
        label, _, _ = CALLOUTS.get(kind, CALLOUTS["NOTE"])
        style = f"Callout {kind if kind in CALLOUTS else 'NOTE'}"
        # split into paragraphs on blank lines; '- ' lines become bullet-like lines
        blocks, cur = [], []
        for ln in lines:
            if not ln.strip():
                if cur:
                    blocks.append(" ".join(cur))
                    cur = []
            elif re.match(r"^\s*([-*]|\d+\.)\s+", ln):
                if cur:
                    blocks.append(" ".join(cur))
                    cur = []
                blocks.append("• " + re.sub(r"^\s*([-*]|\d+\.)\s+", "", ln) if ln.strip()[0] in "-*"
                              else ln.strip())
            else:
                cur.append(ln.strip())
        if cur:
            blocks.append(" ".join(cur))
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = bool(blocks)
        r = p.add_run(label.upper() + (f": {title}" if title else ""))
        r.bold = True
        r.font.name = HEAD_FONT
        r.font.size = Pt(9.5)
        for k, b in enumerate(blocks):
            q = self.doc.add_paragraph(style=style)
            q.paragraph_format.keep_with_next = k < len(blocks) - 1
            add_inline(q, b, size=10)
            self.text_refs += " " + b
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        sp.paragraph_format.line_spacing = 0.6

    def quote(self, lines: list[str]) -> None:
        p = self.doc.add_paragraph(style="Quote Text")
        p.paragraph_format.left_indent = Cm(0.8)
        add_inline(p, " ".join(l.strip() for l in lines), italic=True)

    def figure(self, caption: str, path: str, width_cm: float | None) -> None:
        img = (ROOT / path) if not Path(path).is_absolute() else Path(path)
        if not img.exists():
            img = (self.src_dir / path).resolve()
        if not img.exists():
            raise FileNotFoundError(f"Figure file not found: {path}")
        p = self.doc.add_paragraph(style="Figure")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(img), width=Cm(width_cm or TEXT_WIDTH_CM))
        m = re.match(r"^(Figure\s+[A-Z0-9]+\.\d+)\.?\s*(.*)$", caption)
        cap = self.doc.add_paragraph(style="Figure Caption")
        if m:
            r = cap.add_run(m.group(1) + ". ")
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(INK)
            add_inline(cap, m.group(2), size=9)
            self.figures.append(m.group(1))
        else:
            add_inline(cap, caption, size=9)
            self.warnings.append(f"Figure caption without number: {caption[:50]}")

    def table(self, rows: list[list[str]], aligns: list[str], caption: str | None) -> None:
        if caption:
            p = self.doc.add_paragraph(style="Table Caption")
            m = re.match(r"^(Table\s+[A-Z0-9]+\.\d+)\.?\s*(.*)$", caption)
            if m:
                r = p.add_run(m.group(1) + ". ")
                r.bold = True
                add_inline(p, m.group(2), size=9)
            else:
                add_inline(p, caption, size=9)
        ncols = len(rows[0])
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        # Column widths: every column is at least as wide as its longest word, and the
        # remaining width goes to columns with more text. 0.2 cm ~ one character at 9.5 pt.
        def _clean(x: str) -> str:
            return re.sub(r"[*`]", "", x)

        char_cm, pad_cm = 0.2, 0.45
        mins, prefs = [], []
        for c in range(ncols):
            cells = [_clean(r[c]) if c < len(r) else "" for r in rows]
            longest_word = max((len(w) for cell in cells for w in cell.split()), default=3)
            avg = sum(len(x) for x in cells) / len(cells)
            mins.append(longest_word * char_cm + pad_cm)
            prefs.append(max(mins[-1], min(avg, 80) * char_cm + pad_cm))
        if sum(prefs) <= TEXT_WIDTH_CM:
            cm_widths = [w * TEXT_WIDTH_CM / sum(prefs) for w in prefs]
        else:
            spare = max(TEXT_WIDTH_CM - sum(mins), 0)
            extra = [pf - mn for pf, mn in zip(prefs, mins)]
            cm_widths = [mn + (spare * e / sum(extra) if sum(extra) else 0) for mn, e in zip(mins, extra)]
            scale = TEXT_WIDTH_CM / sum(cm_widths)
            cm_widths = [w * scale for w in cm_widths]
        widths = [Cm(w) for w in cm_widths]
        for j, col in enumerate(t.columns):
            col.width = widths[j]  # sets <w:gridCol>, which LibreOffice and Word both honour
        tbl_pr = t._tbl.tblPr  # t.autofit = False above already set <w:tblLayout w:type="fixed"/>
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), HAIRLINE)
            borders.append(el)
        tbl_pr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)
        for i, row in enumerate(rows):
            tr = t.rows[i]
            if i == 0:  # repeat header row on each page
                trpr = tr._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                trpr.append(th)
            cant = OxmlElement("w:cantSplit")
            tr._tr.get_or_add_trPr().append(cant)
            for j in range(ncols):
                cell = tr.cells[j]
                cell.width = widths[j]
                txt = row[j] if j < len(row) else ""
                para = cell.paragraphs[0]
                para.style = self.doc.styles["Table Text"]
                if aligns and j < len(aligns):
                    para.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}.get(
                        aligns[j], WD_ALIGN_PARAGRAPH.LEFT)
                for k, chunk in enumerate(txt.split("<br>")):
                    if k:
                        para.add_run().add_break(WD_BREAK.LINE)
                    add_inline(para, chunk.strip(), bold=(i == 0), size=9.5, code_size=8.5)
                if i == 0:
                    _set_cell_shading(cell, TABLE_HEAD_FILL)
                self.text_refs += " " + txt
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Markdown block parser
# ---------------------------------------------------------------------------
LIST_RE = re.compile(r"^(?P<indent>\s*)(?P<marker>[-*]|\d+\.)\s+(?P<text>.*)$")
CHECK_RE = re.compile(r"^\s*[-*]\s+\[(?P<mark>[ xX])\]\s+(?P<text>.*)$")
FIG_RE = re.compile(r"^!\[(?P<cap>[^\]]*)\]\((?P<path>[^)]+)\)(\{width=(?P<w>[\d.]+)\})?\s*$")
INCLUDE_RE = re.compile(r"^@include\(\s*(?P<path>[^,)]+)\s*(,\s*(?P<lang>\w+))?\s*\)\s*$")


def parse_front_matter(lines: list[str]) -> tuple[dict, list[str]]:
    if lines and lines[0].strip() == "---":
        meta = {}
        for i in range(1, len(lines)):
            if lines[i].strip() == "---":
                return meta, lines[i + 1:]
            k, _, v = lines[i].partition(":")
            meta[k.strip()] = v.strip()
    return {}, lines


def render_markdown(b: Builder, lines: list[str]) -> None:
    i = 0
    n = len(lines)
    pending_table_caption = None
    while i < n:
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1
            continue
        if s.startswith("<!--"):
            while i < n and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        if re.fullmatch(r"-{3,}|\*{3,}", s):
            b.rule()
            i += 1
            continue
        if s == r"\pagebreak":
            b.page_break()
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            b.heading(len(m.group(1)), m.group(2).strip())
            i += 1
            continue
        if s.startswith("```"):
            info = s[3:].strip()
            lang, _, rest = info.partition(" ")
            body = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                body.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            if lang == "prompt":
                b.prompt(body, rest.strip() or None)
            else:
                tm = re.search(r"title=(\S+)", rest)
                b.code(body, lang, tm.group(1) if tm else None)
            continue
        m = INCLUDE_RE.match(s)
        if m:
            path = ROOT / m.group("path").strip()
            code_lines = path.read_text(encoding="utf-8").rstrip("\n").split("\n")
            b.code(code_lines, m.group("lang") or path.suffix.lstrip("."), m.group("path").strip())
            i += 1
            continue
        m = FIG_RE.match(s)
        if m:
            b.figure(m.group("cap"), m.group("path"), float(m.group("w")) if m.group("w") else None)
            i += 1
            continue
        if s.startswith(">"):
            block = []
            while i < n and lines[i].strip().startswith(">"):
                block.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            cm = re.match(r"^\[!(\w+)\]\s*(.*)$", block[0].strip())
            if cm:
                b.callout(cm.group(1), cm.group(2).strip() or None, block[1:])
            else:
                b.quote(block)
            continue
        if s.startswith("Table ") and re.match(r"^Table\s+[A-Z0-9]+\.\d+", s) and i + 1 < n \
                and lines[i + 1].strip().startswith("|"):
            pending_table_caption = s
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            aligns = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                    aligns = ["c" if c.startswith(":") and c.endswith(":") else
                              "r" if c.endswith(":") else "l" for c in cells]
                else:
                    rows.append(cells)
                i += 1
            b.table(rows, aligns, pending_table_caption)
            pending_table_caption = None
            continue
        cm = CHECK_RE.match(line)
        if cm:
            while i < n and CHECK_RE.match(lines[i]):
                c = CHECK_RE.match(lines[i])
                b.checklist(c.group("text"), c.group("mark") in "xX")
                i += 1
            continue
        lm = LIST_RE.match(line)
        if lm:
            num_id = None
            while i < n:
                lm = LIST_RE.match(lines[i])
                if not lm:
                    # continuation line of the previous item
                    if lines[i].strip() and lines[i].startswith("  ") and not lines[i].strip().startswith(("|", "```", ">")):
                        b.text_refs += " " + lines[i]
                        last = b.doc.paragraphs[-1]
                        add_inline(last, " " + lines[i].strip())
                        i += 1
                        continue
                    break
                if CHECK_RE.match(lines[i]):
                    break
                level = 1 if len(lm.group("indent")) >= 2 else 0
                ordered = lm.group("marker").endswith(".")
                if ordered and level == 0 and num_id is None:
                    num_id = b._new_list_num("List Number")
                b.bullet(lm.group("text"), level, ordered, num_id)
                i += 1
            continue
        # plain paragraph: join following non-blank lines
        para = [s]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,4}\s|```|>|\||!\[|@include|\\pagebreak)",
                                                          lines[i].strip()) \
                and not LIST_RE.match(lines[i]):
            para.append(lines[i].strip())
            i += 1
        b.paragraph(" ".join(para))


def check_figures(b: Builder) -> list[str]:
    """Every numbered figure must be referenced in the text at least once."""
    problems = []
    for fig in b.figures:
        if b.text_refs.count(fig) < 1:
            problems.append(f"{fig} has a caption but is never referenced in the text")
    return problems


def build(sources: list[Path], out: Path, pdf: bool = False) -> Path:
    doc = Document()
    setup_document(doc)
    b = Builder(doc, sources[0].parent)
    header = BOOK_SHORT_TITLE
    for k, src in enumerate(sources):
        lines = src.read_text(encoding="utf-8").split("\n")
        meta, body = parse_front_matter(lines)
        if meta and k == 0:
            b.title_page(meta)
            header = meta.get("header", header)
        render_markdown(b, body)
    setup_header_footer(doc.sections[0], header)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    problems = check_figures(b) + b.warnings
    for p in problems:
        print("WARNING:", p, file=sys.stderr)
    print(f"Built {out} ({len(b.figures)} figures)")
    if pdf:
        pdf_path = out.with_suffix(".pdf")
        pdf_path.unlink(missing_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out.parent.resolve()),
                        str(out.resolve())], check=True, capture_output=True)
        if not pdf_path.exists():  # soffice exits 0 even when it cannot load the file
            raise RuntimeError(f"PDF export failed for {out}")
        print(f"Exported {pdf_path}")
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sources", nargs="+", type=Path)
    ap.add_argument("-o", "--out", type=Path, required=True)
    ap.add_argument("--pdf", action="store_true", help="also export a PDF via LibreOffice")
    args = ap.parse_args()
    build(args.sources, args.out, args.pdf)


if __name__ == "__main__":
    main()
